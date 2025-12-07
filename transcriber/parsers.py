import re
import logging
from pathlib import Path
from typing import List, Dict, Any, Union

log = logging.getLogger(__name__)

class TextParser:
    """Handles extraction and parsing of various text file formats."""

    @staticmethod
    def parse(file_path: Path) -> List[Dict[str, Any]]:
        """Main entry point to parse a file into segments."""
        ext = file_path.suffix.lower()
        content = ""

        if ext == '.pdf':
            content = TextParser._extract_pdf(file_path)
        elif ext == '.docx':
            content = TextParser._extract_docx(file_path)
        elif ext in ('.html', '.htm'):
            content = TextParser._extract_html(file_path)
        else:
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
            except Exception as e:
                log.error(f"Failed to read file: {e}")
                return []

        if ext == '.srt':
            return TextParser._parse_srt(content)
        elif ext == '.vtt':
            return TextParser._parse_vtt(content)
        
        return TextParser._parse_plain_text(content)

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        # Try PyMuPDF
        try:
            import fitz
            with fitz.open(str(path)) as doc:
                return '\n\n'.join(page.get_text() for page in doc)
        except ImportError: pass
        
        # Try pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                return '\n\n'.join(p.extract_text() for p in pdf.pages if p.extract_text())
        except ImportError: pass

        # Try PyPDF2
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(str(path))
            return '\n\n'.join(p.extract_text() for p in reader.pages if p.extract_text())
        except ImportError: pass

        log.error("No suitable PDF library found (pymupdf, pdfplumber, or PyPDF2).")
        return ""

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            from docx import Document
            doc = Document(str(path))
            return '\n\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            log.error("python-docx not installed.")
        except Exception as e:
            log.warning(f"DOCX extraction failed: {e}")
        return ""

    @staticmethod
    def _extract_html(path: Path) -> str:
        try:
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        except Exception:
            return ""

        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(content, 'html.parser')
            
            # Cleaning
            for tag in ['script', 'style', 'nav', 'header', 'footer', 'aside', 'noscript', 'iframe']:
                for el in soup.find_all(tag):
                    el.decompose()
            
            # Content heuristics
            main_content = None
            selectors = ['article', '[role="article"]', 'main', '[role="main"]', '.content', '#content']
            for sel in selectors:
                found = soup.select_one(sel)
                if found and len(found.get_text(strip=True)) > 200:
                    main_content = found
                    break
            
            target = main_content or soup.body or soup
            
            parts = []
            for elem in target.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li', 'blockquote']):
                text = elem.get_text(separator=' ', strip=True)
                if text and len(text) > 20:
                    parts.append(text)
            
            return '\n\n'.join(parts)
            
        except ImportError:
            log.warning("BeautifulSoup4 not installed. HTML parsing will be poor.")
            # Fallback regex parsing
            text = re.sub(r'<[^>]+>', ' ', content)
            return '\n'.join([line.strip() for line in text.split('\n') if line.strip()])

    @staticmethod
    def _parse_srt(content: str) -> List[Dict]:
        segments = []
        for block in re.split(r'\n\n+', content.strip()):
            lines = block.strip().split('\n')
            if len(lines) < 3: continue
            
            # Parse timestamp: 00:00:00,000 --> 00:00:00,000
            m = re.search(r'(\d{2}):(\d{2}):(\d{2}),(\d{3})\s*-->\s*(\d{2}):(\d{2}):(\d{2}),(\d{3})', lines[1])
            if not m: continue
            
            h1, m1, s1, ms1, h2, m2, s2, ms2 = map(int, m.groups())
            start = h1*3600 + m1*60 + s1 + ms1/1000
            end = h2*3600 + m2*60 + s2 + ms2/1000
            
            segments.append({
                'start': start, 'end': end, 
                'text': ' '.join(lines[2:]), 'words': []
            })
        return segments

    @staticmethod
    def _parse_vtt(content: str) -> List[Dict]:
        segments = []
        lines = content.strip().split('\n')
        i = 0
        while i < len(lines):
            line = lines[i]
            if '-->' in line:
                # Support both 00:00.000 and 00:00:00.000 formats
                m = re.search(r'(?:(\d{2}):)?(\d{2}):(\d{2})\.(\d{3})\s*-->\s*(?:(\d{2}):)?(\d{2}):(\d{2})\.(\d{3})', line)
                if m:
                    # Logic to handle optional hours
                    groups = list(m.groups())
                    if groups[0] is None: groups[0] = '00'
                    if groups[4] is None: groups[4] = '00'
                    
                    h1, m1, s1, ms1, h2, m2, s2, ms2 = map(int, groups)
                    start = h1*3600 + m1*60 + s1 + ms1/1000
                    end = h2*3600 + m2*60 + s2 + ms2/1000
                    
                    text_lines = []
                    i += 1
                    while i < len(lines) and lines[i].strip() and '-->' not in lines[i]:
                        text_lines.append(lines[i].strip())
                        i += 1
                    
                    segments.append({
                        'start': start, 'end': end,
                        'text': ' '.join(text_lines), 'words': []
                    })
                    continue
            i += 1
        return segments

    @staticmethod
    def _parse_plain_text(content: str) -> List[Dict]:
        """Splits plain text into readable sentence-sized chunks."""
        sentences = TextParser._split_sentences(content)
        segments = []
        current_chunk = []
        current_word_count = 0
        
        target_words = 12
        max_words = 20

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence: continue
            
            words = sentence.split()
            wc = len(words)
            
            # Handle very long sentences
            if wc > max_words:
                if current_chunk:
                    segments.append(' '.join(current_chunk))
                    current_chunk, current_word_count = [], 0
                
                # Split the long sentence itself
                segments.extend(TextParser._split_long_sentence(sentence, target_words))
                continue

            if current_word_count + wc > max_words and current_word_count > 5:
                segments.append(' '.join(current_chunk))
                current_chunk, current_word_count = [], 0
            
            current_chunk.append(sentence)
            current_word_count += wc
            
            if current_word_count >= target_words:
                segments.append(' '.join(current_chunk))
                current_chunk, current_word_count = [], 0
        
        if current_chunk:
            segments.append(' '.join(current_chunk))
            
        return [{'start': i, 'end': i+1, 'text': t, 'words': []} for i, t in enumerate(segments) if t.strip()]

    @staticmethod
    def _split_sentences(text: str) -> List[str]:
        # Protected abbreviations
        abbreviations = ['Mr.', 'Ms.', 'Dr.', 'Prof.', 'etc.', 'e.g.', 'i.e.']
        protected = text
        placeholders = {}
        
        for i, abbr in enumerate(abbreviations):
            ph = f"__ABBR{i}__"
            if abbr in protected:
                protected = protected.replace(abbr, ph)
                placeholders[ph] = abbr
        
        # Split on punctuation followed by space and uppercase
        pattern = r'(?<=[.!?])\s+(?=[A-ZÄÖÜÁÉÍÓÚÀÈÌÒÙÂÊÎÔÛÅÆØÑ])'
        raw_sents = re.split(pattern, protected)
        
        final_sents = []
        for s in raw_sents:
            for ph, abbr in placeholders.items():
                s = s.replace(ph, abbr)
            final_sents.append(s)
        return final_sents

    @staticmethod
    def _split_long_sentence(sentence: str, target: int) -> List[str]:
        words = sentence.split()
        if len(words) <= target + 5: return [sentence]
        
        chunks = []
        current = []
        for i, w in enumerate(words):
            current.append(w)
            if len(current) >= target and (w.endswith((',', ';', ':')) or len(current) > target+5):
                chunks.append(' '.join(current))
                current = []
        if current:
            if chunks and len(current) < 4:
                chunks[-1] += ' ' + ' '.join(current)
            else:
                chunks.append(' '.join(current))
        return chunks